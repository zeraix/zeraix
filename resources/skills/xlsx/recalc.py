"""Recalculate every formula in a workbook with LibreOffice and report formula errors.

    python recalc.py model.xlsx                 # scan: prints a JSON report, the original file is untouched
    python recalc.py model.xlsx --inplace       # also replace model.xlsx with the recalculated copy (cached values)
    python recalc.py model.xlsx --timeout 300

Why: openpyxl writes formulas as text without results. Excel computes them on open, but nothing has checked
them for #REF! / #DIV/0! / #NAME? / #VALUE! / #N/A until this script runs them through a real engine.
Exit code 0 = no errors, 1 = errors found (listed in the report), 2 = could not run.
"""
from __future__ import annotations

import argparse
import json
import shutil
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "office"))
import lo  # noqa: E402

ERROR_TOKENS = ("#REF!", "#DIV/0!", "#NAME?", "#VALUE!", "#N/A", "#NUM!", "#NULL!", "#SPILL!", "#CALC!", "Err:")


def is_formula(v) -> bool:
    if isinstance(v, str):
        return v.startswith("=")
    return type(v).__name__ in ("ArrayFormula", "DataTableFormula")


def formula_text(v) -> str:
    if isinstance(v, str):
        return v
    return getattr(v, "text", None) or str(v)


def scan(original: Path, recalculated: Path) -> dict:
    from openpyxl import load_workbook

    wb_f = load_workbook(original, data_only=False)
    wb_v = load_workbook(recalculated, data_only=True)
    formulas = 0
    errors = []
    sheets = []
    for ws in wb_f.worksheets:
        vs = wb_v[ws.title] if ws.title in wb_v.sheetnames else None
        count = 0
        for row in ws.iter_rows():
            for cell in row:
                if not is_formula(cell.value):
                    continue
                count += 1
                value = vs[cell.coordinate].value if vs is not None else None
                if isinstance(value, str) and value.startswith(ERROR_TOKENS):
                    errors.append({"sheet": ws.title, "cell": cell.coordinate, "value": value,
                                   "formula": formula_text(cell.value)})
        formulas += count
        sheets.append({"sheet": ws.title, "formulas": count})
    return {"formulas": formulas, "sheets": sheets, "errors": errors}


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("file", help=".xlsx / .xlsm workbook")
    ap.add_argument("--inplace", action="store_true", help="replace the file with the recalculated copy")
    ap.add_argument("--keep", help="also keep the recalculated copy at this path")
    ap.add_argument("--timeout", type=int, default=180)
    a = ap.parse_args()

    src = Path(a.file).resolve()
    if not src.exists():
        print(json.dumps({"error": f"no such file: {src}"}))
        return 2
    with tempfile.TemporaryDirectory(prefix="recalc-") as td:
        work = Path(td) / src.name
        shutil.copy2(src, work)
        try:
            out = lo.convert(work, src.suffix.lstrip(".").lower() or "xlsx", Path(td) / "out", timeout=a.timeout)
        except Exception as e:
            print(json.dumps({"error": f"LibreOffice recalculation failed: {e}"}))
            return 2
        report = scan(src, out)
        report["file"] = str(src)
        if a.keep:
            shutil.copy2(out, a.keep)
            report["recalculated_copy"] = str(Path(a.keep).resolve())
        if a.inplace:
            shutil.copy2(out, src)
            report["replaced_in_place"] = True
    report["status"] = "errors" if report["errors"] else "ok"
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 1 if report["errors"] else 0


if __name__ == "__main__":
    sys.exit(main())
